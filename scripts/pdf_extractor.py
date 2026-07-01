"""
pdf_extractor.py

Project:
PDF Extractor

Purpose:
Extract invoice-style data from PDF files and generate a formatted Excel report.

Expected folder structure:

Python/
    datasets/
        pdfs/
            invoice_001.pdf
            invoice_002.pdf
            invoice_003.pdf
    scripts/
        pdf_extractor.py
    outputs/
    archive/

Run from the Python folder:
    py scripts/pdf_extractor.py

Install packages once:
    py -m pip install pandas openpyxl pdfplumber

Output format:
    outputs/pdf_extraction_report_06.30_1.xlsx
    outputs/pdf_extraction_report_06.30_2.xlsx

Output sheets:
    Summary
    Extracted_Data
    Invalid_PDFs

Invalid_PDFs columns:
    File Name
    Issue

What it extracts:
    - File Name
    - Invoice Number
    - Vendor
    - Invoice Date
    - Amount Due
    - Email
    - Phone
    - Address Candidate
    - Word Count
    - Issues

Notes:
    - This is designed for text-based PDFs.
    - Scanned image PDFs may need OCR first.
"""

import re
from pathlib import Path
from datetime import datetime

import pandas as pd
import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent.parent
DATASETS_DIR = BASE_DIR / "datasets"
PDF_DIR = DATASETS_DIR / "pdfs"
OUTPUTS_DIR = BASE_DIR / "outputs"
ARCHIVE_DIR = BASE_DIR / "archive"

PDF_DIR.mkdir(parents=True, exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)
ARCHIVE_DIR.mkdir(exist_ok=True)


def get_next_output_file(output_dir: Path, base_name: str = "pdf_extraction_report", extension: str = ".xlsx") -> Path:
    today = datetime.now().strftime("%m.%d")
    existing_files = list(output_dir.glob(f"{base_name}_{today}_*{extension}"))

    versions = []
    for file in existing_files:
        match = re.search(rf"^{re.escape(base_name)}_{re.escape(today)}_(\d+)$", file.stem)
        if match:
            versions.append(int(match.group(1)))

    next_version = max(versions, default=0) + 1
    return output_dir / f"{base_name}_{today}_{next_version}{extension}"


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def unique_join(values, limit=5):
    output = []
    seen = set()

    for value in values:
        value = clean_text(value)
        if not value:
            continue

        key = value.lower()
        if key not in seen:
            output.append(value)
            seen.add(key)

    return ", ".join(output[:limit])


def extract_text_from_pdf(pdf_path):
    text_pages = []

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_pages.append(page_text)

    return "\n".join(text_pages)


def extract_first(patterns, text):
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return clean_text(match.group(1))
    return ""


def extract_emails(text):
    return sorted(set(re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)))


def extract_phone_numbers(text):
    pattern = r"(?:\+?\d[\d\s().-]{7,}\d)"
    phones = re.findall(pattern, text)

    cleaned = []
    for phone in phones:
        digits = re.sub(r"\D", "", phone)
        if 8 <= len(digits) <= 15:
            cleaned.append(phone.strip())

    return sorted(set(cleaned))


def extract_address_candidates(text):
    candidates = []

    address_pattern = r"\b\d{2,6}\s+[A-Za-z0-9\s.,#'-]+(?:Street|St\.?|Avenue|Ave\.?|Road|Rd\.?|Boulevard|Blvd\.?|Drive|Dr\.?|Lane|Ln\.?|Court|Ct\.?|Way|Highway|Hwy\.?|Suite|Ste\.?)\s*[A-Za-z0-9\s.,#'-]*"
    candidates += re.findall(address_pattern, text, flags=re.IGNORECASE)

    return sorted(set(candidate.strip() for candidate in candidates))


def parse_amount(value):
    value = clean_text(value)
    if not value:
        return None

    cleaned = value.replace("$", "").replace(",", "").strip()

    try:
        return float(cleaned)
    except ValueError:
        return None


def parse_pdf(pdf_path):
    issues = []

    result = {
        "File Name": pdf_path.name,
        "Invoice Number": "",
        "Vendor": "",
        "Invoice Date": "",
        "Amount Due": "",
        "Email": "",
        "Phone": "",
        "Address Candidate": "",
        "Word Count": 0,
        "Issues": "",
    }

    try:
        text = extract_text_from_pdf(pdf_path)
    except Exception as exc:
        return None, {"File Name": pdf_path.name, "Issue": f"PDF read failed: {exc}"}

    if not text.strip():
        return None, {"File Name": pdf_path.name, "Issue": "No extractable text found. PDF may be scanned and need OCR."}

    result["Word Count"] = len(re.findall(r"\w+", text))

    invoice_number = extract_first([
        r"Invoice\s*(?:Number|No\.?|#)\s*[:\-]?\s*([A-Za-z0-9\-]+)",
        r"INV[-\s]?([0-9]{3,})",
    ], text)

    vendor = extract_first([
        r"Vendor\s*[:\-]?\s*(.+)",
        r"From\s*[:\-]?\s*(.+)",
        r"Company\s*[:\-]?\s*(.+)",
    ], text)

    invoice_date = extract_first([
        r"Invoice\s*Date\s*[:\-]?\s*([A-Za-z0-9,\/\-\s]+)",
        r"Date\s*[:\-]?\s*([A-Za-z0-9,\/\-\s]+)",
    ], text)

    amount_raw = extract_first([
        r"Amount\s*Due\s*[:\-]?\s*\$?\s*([0-9,]+(?:\.[0-9]{2})?)",
        r"Total\s*Due\s*[:\-]?\s*\$?\s*([0-9,]+(?:\.[0-9]{2})?)",
        r"Total\s*[:\-]?\s*\$?\s*([0-9,]+(?:\.[0-9]{2})?)",
    ], text)

    amount_due = parse_amount(amount_raw)

    emails = extract_emails(text)
    phones = extract_phone_numbers(text)
    addresses = extract_address_candidates(text)

    result["Invoice Number"] = invoice_number
    result["Vendor"] = vendor
    result["Invoice Date"] = invoice_date
    result["Amount Due"] = amount_due
    result["Email"] = unique_join(emails, limit=3)
    result["Phone"] = unique_join(phones, limit=3)
    result["Address Candidate"] = unique_join(addresses, limit=3)

    if not result["Invoice Number"]:
        issues.append("Missing invoice number")

    if not result["Vendor"]:
        issues.append("Missing vendor")

    if not result["Invoice Date"]:
        issues.append("Missing invoice date")

    if result["Amount Due"] is None:
        issues.append("Missing or invalid amount due")

    if not result["Email"]:
        issues.append("Missing email")

    if not result["Phone"]:
        issues.append("Missing phone")

    if not result["Address Candidate"]:
        issues.append("Missing address candidate")

    if result["Word Count"] < 20:
        issues.append("Low text extraction count")

    result["Issues"] = "; ".join(issues)

    return result, None


def build_summary(extracted_df, invalid_df):
    total_amount = extracted_df["Amount Due"].fillna(0).sum() if len(extracted_df) else 0
    rows_with_issues = extracted_df["Issues"].astype(str).str.strip().ne("").sum() if len(extracted_df) else 0

    return pd.DataFrame([
        ["PDF folder", str(PDF_DIR)],
        ["PDFs processed", len(extracted_df) + len(invalid_df)],
        ["PDFs extracted", len(extracted_df)],
        ["Invalid/unreadable PDFs", len(invalid_df)],
        ["Extracted rows with issues", int(rows_with_issues)],
        ["Total extracted amount due", total_amount],
        ["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ], columns=["Metric", "Value"])


def get_col(ws, name):
    for cell in ws[1]:
        if cell.value == name:
            return cell.column
    return None


def format_workbook(output_file):
    wb = load_workbook(output_file)

    dark_fill = PatternFill("solid", fgColor="111827")
    green_fill = PatternFill("solid", fgColor="166534")
    red_fill = PatternFill("solid", fgColor="991B1B")
    summary_fill = PatternFill("solid", fgColor="E0F2FE")
    issue_fill = PatternFill("solid", fgColor="F4CCCC")
    invalid_cell_fill = PatternFill("solid", fgColor="FCE4E4")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    white_font = Font(color="FFFFFF", bold=True)

    thin_gray = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    issue_map = {
        "Missing invoice number": ["Invoice Number"],
        "Missing vendor": ["Vendor"],
        "Missing invoice date": ["Invoice Date"],
        "Missing or invalid amount due": ["Amount Due"],
        "Missing email": ["Email"],
        "Missing phone": ["Phone"],
        "Missing address candidate": ["Address Candidate"],
        "Low text extraction count": ["Word Count"],
    }

    for ws in wb.worksheets:
        max_row = ws.max_row
        max_col = ws.max_column

        if ws.title == "Summary":
            ws.freeze_panes = None
        else:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = dark_fill
            cell.font = white_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        if ws.title == "Extracted_Data":
            for cell in ws[1]:
                cell.fill = green_fill

        if ws.title == "Invalid_PDFs":
            for cell in ws[1]:
                cell.fill = red_fill

        if ws.title == "Summary":
            for row in range(2, max_row + 1):
                ws[f"A{row}"].fill = summary_fill

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        amount_col = get_col(ws, "Amount Due")
        if amount_col:
            for row in range(2, max_row + 1):
                ws.cell(row, amount_col).number_format = "$#,##0.00"

        word_count_col = get_col(ws, "Word Count")
        if word_count_col:
            for row in range(2, max_row + 1):
                ws.cell(row, word_count_col).number_format = "#,##0"

        issues_col = get_col(ws, "Issues")
        if issues_col:
            ws.cell(1, issues_col).fill = PatternFill("solid", fgColor="D99694")
            ws.cell(1, issues_col).font = Font(color="000000", bold=True)

            for row in range(2, max_row + 1):
                issues = str(ws.cell(row, issues_col).value or "")
                if issues:
                    ws.cell(row, issues_col).fill = issue_fill

                    if "Low text extraction count" in issues:
                        for cell in ws[row]:
                            if cell.fill.fill_type is None:
                                cell.fill = warning_fill
                        ws.cell(row, issues_col).fill = issue_fill

                    for issue, columns in issue_map.items():
                        if issue in issues:
                            for col_name in columns:
                                col = get_col(ws, col_name)
                                if col:
                                    ws.cell(row, col).fill = invalid_cell_fill

        if max_row >= 2 and max_col >= 1 and ws.title != "Summary":
            table_name = re.sub(r"[^A-Za-z0-9_]", "", ws.title)[:20] + "Table"
            ref = f"A1:{get_column_letter(max_col)}{max_row}"
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            try:
                ws.add_table(table)
            except ValueError:
                pass

        for col_idx in range(1, max_col + 1):
            header = str(ws.cell(1, col_idx).value or "")
            width = 18

            if header in ["File Name", "Invoice Number", "Invoice Date"]:
                width = 24
            elif header in ["Vendor", "Email", "Phone"]:
                width = 30
            elif header in ["Address Candidate", "Issues", "Issue"]:
                width = 48

            ws.column_dimensions[get_column_letter(col_idx)].width = width

        for row in range(1, max_row + 1):
            ws.row_dimensions[row].height = 24

    wb.save(output_file)


def main():
    output_file = get_next_output_file(OUTPUTS_DIR)

    pdf_files = sorted(PDF_DIR.glob("*.pdf"))

    if not pdf_files:
        raise FileNotFoundError(
            f"No PDF files found in:\n{PDF_DIR}\n\n"
            "Place your PDF files inside datasets/pdfs/."
        )

    extracted_rows = []
    invalid_rows = []

    print("Extracting PDF data...")
    for pdf_file in pdf_files:
        print(f"Reading: {pdf_file.name}")
        result, invalid = parse_pdf(pdf_file)

        if invalid:
            invalid_rows.append(invalid)
        else:
            extracted_rows.append(result)

    extracted_df = pd.DataFrame(extracted_rows)
    invalid_df = pd.DataFrame(invalid_rows, columns=["File Name", "Issue"])
    summary_df = build_summary(extracted_df, invalid_df)

    print("Writing extraction report...")
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        extracted_df.to_excel(writer, sheet_name="Extracted_Data", index=False)
        invalid_df.to_excel(writer, sheet_name="Invalid_PDFs", index=False)

    print("Applying Excel formatting...")
    format_workbook(output_file)

    print("\nPDF extraction complete.")
    print(f"Created: {output_file.name}")
    print(f"Saved in: {OUTPUTS_DIR}")
    print(f"PDFs extracted: {len(extracted_df)}")
    print(f"Invalid/unreadable PDFs: {len(invalid_df)}")


if __name__ == "__main__":
    main()
