"""
prestashop_validator_formatted.py

Project:
PrestaShop Validator

Purpose:
Validate a PrestaShop-style product export, split valid and invalid rows,
and generate a formatted Excel report with highlighted issue cells.

Expected folder structure:

Python/
    datasets/
        prestashop_export_messy.xlsx
    scripts/
        prestashop_validator_formatted.py
    outputs/
    archive/

Run from the Python folder:
    py scripts/prestashop_validator_formatted.py

Output format:
    outputs/prestashop_validation_06.30_1.xlsx
    outputs/prestashop_validation_06.30_2.xlsx

Output sheets:
    Summary
    PrestaShop_Valid
    PrestaShop_Invalid

Main features:
    - Cleans product IDs, references, names, categories, prices, quantities, active values, URLs, SEO fields, EAN13, and dates
    - Validates common PrestaShop product export issues
    - Moves invalid rows into PrestaShop_Invalid
    - Adds final Issues column to PrestaShop_Invalid
    - Highlights issue-related cells
    - Sorts invalid rows by Reference so duplicates appear consecutively
    - Formats workbook for portfolio/client presentation
"""

import re
from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent.parent
DATASETS_DIR = BASE_DIR / "datasets"
OUTPUTS_DIR = BASE_DIR / "outputs"
ARCHIVE_DIR = BASE_DIR / "archive"

INPUT_FILE = DATASETS_DIR / "prestashop_export_messy.xlsx"

DATASETS_DIR.mkdir(exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)
ARCHIVE_DIR.mkdir(exist_ok=True)


def get_next_output_file(output_dir: Path, base_name: str = "prestashop_validation", extension: str = ".xlsx") -> Path:
    today = datetime.now().strftime("%m.%d")
    existing_files = list(output_dir.glob(f"{base_name}_{today}_*{extension}"))

    versions = []
    for file in existing_files:
        match = re.search(rf"^{re.escape(base_name)}_{re.escape(today)}_(\d+)$", file.stem)
        if match:
            versions.append(int(match.group(1)))

    next_version = max(versions, default=0) + 1
    return output_dir / f"{base_name}_{today}_{next_version}{extension}"


def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def clean_int(value):
    if pd.isna(value) or str(value).strip() == "":
        return None
    try:
        return int(value)
    except ValueError:
        return None


def clean_float(value):
    if pd.isna(value) or str(value).strip() == "":
        return None
    cleaned = str(value).replace(",", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def normalize_category(value):
    return clean_text(value).title()


def normalize_active(value):
    value = clean_text(value).lower()
    active_map = {
        "1": "1",
        "yes": "1",
        "active": "1",
        "true": "1",
        "0": "0",
        "no": "0",
        "inactive": "0",
        "false": "0",
    }
    return active_map.get(value, value)


def parse_date(value):
    if pd.isna(value) or str(value).strip() == "":
        return ""

    parsed = pd.to_datetime(value, errors="coerce", dayfirst=False)

    if pd.isna(parsed):
        parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)

    if pd.isna(parsed):
        return ""

    return parsed.strftime("%Y-%m-%d")


def is_valid_url(value):
    value = clean_text(value)
    return value.startswith("http://") or value.startswith("https://")


def is_valid_ean13(value):
    value = clean_text(value)
    if not value:
        return True
    return bool(re.fullmatch(r"\d{13}", value))


def add_issue(issues_by_index, row_index, issue):
    if row_index not in issues_by_index:
        issues_by_index[row_index] = []
    issues_by_index[row_index].append(issue)


def count_total_issues(issues_by_index):
    return sum(len(issues) for issues in issues_by_index.values())


def split_valid_invalid(df, issues_by_index):
    invalid_indexes = sorted(issues_by_index.keys())

    valid_df = df.drop(index=invalid_indexes, errors="ignore").copy()
    invalid_df = df.loc[df.index.intersection(invalid_indexes)].copy()

    if len(invalid_df) > 0:
        invalid_df["Issues"] = invalid_df.index.map(
            lambda idx: "; ".join(sorted(set(issues_by_index.get(idx, []))))
        )

        invalid_df["_original_index"] = invalid_df.index
        invalid_df = invalid_df.sort_values(
            by=["Reference", "_original_index"],
            kind="stable"
        ).drop(columns=["_original_index"])
    else:
        invalid_df["Issues"] = []

    return valid_df.reset_index(drop=True), invalid_df.reset_index(drop=True)


def validate_prestashop(df):
    issues_by_index = {}

    df = df.copy()
    df.columns = [str(col).strip() for col in df.columns]

    required_columns = [
        "Product ID", "Reference", "Name", "Category", "Price Tax Excl",
        "Tax Rule", "Quantity", "Active", "Image URL", "Product URL",
        "Meta Title", "Meta Description", "EAN13", "Last Updated"
    ]

    missing_columns = [col for col in required_columns if col not in df.columns]
    if missing_columns:
        raise ValueError(f"Missing required columns: {missing_columns}")

    df["Product ID"] = df["Product ID"].apply(clean_int)
    df["Reference"] = df["Reference"].apply(clean_text)
    df["Name"] = df["Name"].apply(lambda x: clean_text(x).title())
    df["Category"] = df["Category"].apply(normalize_category)
    df["Price Tax Excl"] = df["Price Tax Excl"].apply(clean_float)
    df["Tax Rule"] = df["Tax Rule"].apply(clean_text)
    df["Quantity"] = df["Quantity"].apply(clean_int)
    df["Active"] = df["Active"].apply(normalize_active)
    df["Image URL"] = df["Image URL"].apply(clean_text)
    df["Product URL"] = df["Product URL"].apply(clean_text)
    df["Meta Title"] = df["Meta Title"].apply(clean_text)
    df["Meta Description"] = df["Meta Description"].apply(clean_text)
    df["EAN13"] = df["EAN13"].apply(clean_text)
    df["Last Updated"] = df["Last Updated"].apply(parse_date)

    allowed_tax_rules = {"US Standard", "EU VAT", "No Tax"}

    for index, row in df.iterrows():
        if pd.isna(row["Product ID"]):
            add_issue(issues_by_index, index, "Missing or invalid Product ID")

        if not row["Reference"]:
            add_issue(issues_by_index, index, "Missing Reference/SKU")

        if not row["Name"]:
            add_issue(issues_by_index, index, "Missing product name")

        if not row["Category"]:
            add_issue(issues_by_index, index, "Missing category")

        if pd.isna(row["Price Tax Excl"]):
            add_issue(issues_by_index, index, "Missing or invalid price")
        elif row["Price Tax Excl"] <= 0:
            add_issue(issues_by_index, index, "Price is zero or negative")

        if not row["Tax Rule"]:
            add_issue(issues_by_index, index, "Missing tax rule")
        elif row["Tax Rule"] not in allowed_tax_rules:
            add_issue(issues_by_index, index, "Invalid tax rule")

        if pd.isna(row["Quantity"]):
            add_issue(issues_by_index, index, "Missing or invalid quantity")
        elif row["Quantity"] < 0:
            add_issue(issues_by_index, index, "Quantity is negative")

        if row["Active"] not in ["0", "1"]:
            add_issue(issues_by_index, index, "Invalid active value")

        if not row["Image URL"]:
            add_issue(issues_by_index, index, "Missing image URL")
        elif not is_valid_url(row["Image URL"]):
            add_issue(issues_by_index, index, "Invalid image URL")

        if not row["Product URL"]:
            add_issue(issues_by_index, index, "Missing product URL")
        elif not is_valid_url(row["Product URL"]):
            add_issue(issues_by_index, index, "Invalid product URL")

        if not row["Meta Title"]:
            add_issue(issues_by_index, index, "Missing meta title")

        if not row["Meta Description"]:
            add_issue(issues_by_index, index, "Missing meta description")
        elif len(row["Meta Description"]) < 30:
            add_issue(issues_by_index, index, "Meta description too short")

        if not is_valid_ean13(row["EAN13"]):
            add_issue(issues_by_index, index, "Invalid EAN13")

        if not row["Last Updated"]:
            add_issue(issues_by_index, index, "Missing or invalid last updated date")

    # Duplicate Reference/SKU rule:
    # If a Reference appears more than once, ALL matching rows are invalid.
    duplicate_mask = df.duplicated(subset=["Reference"], keep=False) & df["Reference"].ne("")
    for index in df[duplicate_mask].index:
        add_issue(issues_by_index, index, "Duplicate Reference/SKU")

    valid_df, invalid_df = split_valid_invalid(df, issues_by_index)

    return valid_df, invalid_df, issues_by_index


def get_column_index(ws, column_name):
    for cell in ws[1]:
        if cell.value == column_name:
            return cell.column
    return None


def format_workbook(output_file):
    wb = load_workbook(output_file)

    dark_header_fill = PatternFill("solid", fgColor="1F2937")
    dark_header_font = Font(color="FFFFFF", bold=True)
    valid_header_fill = PatternFill("solid", fgColor="166534")
    invalid_header_fill = PatternFill("solid", fgColor="991B1B")
    issue_header_fill = PatternFill("solid", fgColor="D99694")
    issue_fill = PatternFill("solid", fgColor="F4CCCC")
    invalid_cell_fill = PatternFill("solid", fgColor="FCE4E4")
    duplicate_fill = PatternFill("solid", fgColor="FCE5CD")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    summary_fill = PatternFill("solid", fgColor="E0F2FE")

    thin_gray = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    issue_to_columns = {
        "Missing or invalid Product ID": ["Product ID"],
        "Missing Reference/SKU": ["Reference"],
        "Duplicate Reference/SKU": ["Reference"],
        "Missing product name": ["Name"],
        "Missing category": ["Category"],
        "Missing or invalid price": ["Price Tax Excl"],
        "Price is zero or negative": ["Price Tax Excl"],
        "Missing tax rule": ["Tax Rule"],
        "Invalid tax rule": ["Tax Rule"],
        "Missing or invalid quantity": ["Quantity"],
        "Quantity is negative": ["Quantity"],
        "Invalid active value": ["Active"],
        "Missing image URL": ["Image URL"],
        "Invalid image URL": ["Image URL"],
        "Missing product URL": ["Product URL"],
        "Invalid product URL": ["Product URL"],
        "Missing meta title": ["Meta Title"],
        "Missing meta description": ["Meta Description"],
        "Meta description too short": ["Meta Description"],
        "Invalid EAN13": ["EAN13"],
        "Missing or invalid last updated date": ["Last Updated"],
    }

    for ws in wb.worksheets:
        max_row = ws.max_row
        max_col = ws.max_column

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = dark_header_fill
            cell.font = dark_header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        if "Valid" in ws.title:
            for cell in ws[1]:
                cell.fill = valid_header_fill

        if "Invalid" in ws.title:
            for cell in ws[1]:
                cell.fill = invalid_header_fill

        if ws.title == "Summary":
            ws.freeze_panes = None

            for row in ws.iter_rows():
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

            ws["A1"].fill = dark_header_fill
            ws["A1"].font = dark_header_font
            ws["B1"].fill = dark_header_fill
            ws["B1"].font = dark_header_font

            for row in range(2, max_row + 1):
                ws[f"A{row}"].fill = summary_fill

            ws.column_dimensions["A"].width = 34
            ws.column_dimensions["B"].width = 95
            continue

        for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        # Number formats
        price_col = get_column_index(ws, "Price Tax Excl")
        if price_col:
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=price_col).number_format = "#,##0.00"

        quantity_col = get_column_index(ws, "Quantity")
        if quantity_col:
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=quantity_col).number_format = "0"

        date_col = get_column_index(ws, "Last Updated")
        if date_col:
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=date_col).number_format = "yyyy-mm-dd"

        # Invalid sheet issue highlighting
        if "Invalid" in ws.title:
            issues_col = get_column_index(ws, "Issues")

            if issues_col:
                ws.cell(row=1, column=issues_col).fill = issue_header_fill
                ws.cell(row=1, column=issues_col).font = Font(color="000000", bold=True)

                for row in range(2, max_row + 1):
                    issues_text = str(ws.cell(row=row, column=issues_col).value or "")
                    ws.cell(row=row, column=issues_col).fill = issue_fill

                    if "Duplicate Reference/SKU" in issues_text:
                        for cell in ws[row]:
                            cell.fill = duplicate_fill
                        ws.cell(row=row, column=issues_col).fill = issue_fill

                    if "Meta description too short" in issues_text:
                        for cell in ws[row]:
                            cell.fill = warning_fill
                        ws.cell(row=row, column=issues_col).fill = issue_fill

                    for issue, columns in issue_to_columns.items():
                        if issue in issues_text:
                            for col_name in columns:
                                col_idx = get_column_index(ws, col_name)
                                if col_idx:
                                    ws.cell(row=row, column=col_idx).fill = invalid_cell_fill

        # Create Excel table
        if max_row >= 2 and max_col >= 1:
            table_name = re.sub(r"[^A-Za-z0-9_]", "", ws.title)[:20] + "Table"
            ref = f"A1:{get_column_letter(max_col)}{max_row}"
            table = Table(displayName=table_name, ref=ref)
            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            try:
                ws.add_table(table)
            except ValueError:
                pass

        # Column widths
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            header = str(ws.cell(row=1, column=col_idx).value or "")

            if header == "Issues":
                ws.column_dimensions[col_letter].width = 62
            elif header in ["Image URL", "Product URL", "Meta Description"]:
                ws.column_dimensions[col_letter].width = 44
            elif header in ["Name", "Meta Title"]:
                ws.column_dimensions[col_letter].width = 28
            elif header in ["Reference", "Category", "Tax Rule", "Last Updated", "EAN13"]:
                ws.column_dimensions[col_letter].width = 18
            else:
                ws.column_dimensions[col_letter].width = 14

        for row in range(1, max_row + 1):
            ws.row_dimensions[row].height = 22

    wb.save(output_file)


def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Input file not found:\n{INPUT_FILE}\n\n"
            "Place prestashop_export_messy.xlsx inside the datasets folder."
        )

    output_file = get_next_output_file(OUTPUTS_DIR)

    print("Reading PrestaShop export...")
    print(f"Input: {INPUT_FILE}")

    df = pd.read_excel(INPUT_FILE, sheet_name="PrestaShop_Raw")

    print("Validating PrestaShop data...")
    valid_df, invalid_df, issues = validate_prestashop(df)

    summary = pd.DataFrame([
        ["Input file", str(INPUT_FILE)],
        ["Output file", str(output_file)],
        ["Raw rows", len(df)],
        ["Valid rows", len(valid_df)],
        ["Invalid rows moved", len(invalid_df)],
        ["Issue count", count_total_issues(issues)],
        ["Duplicate rule", "All rows are invalid when Reference/SKU appears more than once"],
        ["Invalid sorting", "PrestaShop_Invalid is sorted by Reference so duplicates appear consecutively"],
        ["Formatting", "Issue cells and Issues column are highlighted automatically"],
        ["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ], columns=["Metric", "Value"])

    print("Writing output workbook...")
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        valid_df.to_excel(writer, sheet_name="PrestaShop_Valid", index=False)
        invalid_df.to_excel(writer, sheet_name="PrestaShop_Invalid", index=False)

    print("Applying Excel formatting...")
    format_workbook(output_file)

    print("\nPrestaShop validation complete.")
    print(f"Created: {output_file.name}")
    print(f"Saved in: {OUTPUTS_DIR}")
    print(f"Valid rows: {len(valid_df)}")
    print(f"Invalid rows moved: {len(invalid_df)}")
    print(f"Issue count: {count_total_issues(issues)}")


if __name__ == "__main__":
    main()
