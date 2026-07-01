"""
business_website_lead_scraper.py

Project:
Business Website Lead Scraper

Use case:
Dental, construction, local service, agency, clinic, contractor, and business websites.

Purpose:
Read a list of business website URLs, crawl important internal pages, extract useful
business/contact information, and generate a formatted Excel report.

Expected folder structure:

Python/
    datasets/
        business_websites.xlsx
    scripts/
        business_website_lead_scraper.py
    outputs/
    archive/

Required input sheet:
    Websites_Raw

Required input columns:
    URL

Optional input columns:
    Business Type
    Notes

Run from the Python folder:
    py scripts/business_website_lead_scraper.py

Install packages once:
    py -m pip install pandas openpyxl requests beautifulsoup4

Output:
    outputs/business_leads_scrape_06.30_1.xlsx

Output sheets:
    Summary
    Business_Leads
    Invalid_URLs

Invalid_URLs columns:
    URL
    Issue
"""

import re
import time
from pathlib import Path
from datetime import datetime
from urllib.parse import urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent.parent
DATASETS_DIR = BASE_DIR / "datasets"
OUTPUTS_DIR = BASE_DIR / "outputs"
ARCHIVE_DIR = BASE_DIR / "archive"

INPUT_FILE = DATASETS_DIR / "business_websites.xlsx"

DATASETS_DIR.mkdir(exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)
ARCHIVE_DIR.mkdir(exist_ok=True)

MAX_PAGES_PER_SITE = 12
REQUEST_DELAY_SECONDS = 0.5

IMPORTANT_PAGE_KEYWORDS = [
    "contact", "about", "team", "leadership", "staff", "doctor", "dentist",
    "owner", "founder", "management", "services", "service", "location",
    "locations", "office", "company", "who-we-are", "our-story", "meet",
    "crew", "projects", "portfolio"
]

SOCIAL_DOMAINS = [
    "facebook.com", "instagram.com", "linkedin.com", "twitter.com", "x.com",
    "youtube.com", "tiktok.com"
]

SERVICE_KEYWORDS = [
    "general dentistry", "cosmetic dentistry", "dental implants", "teeth whitening",
    "orthodontics", "root canal", "emergency dental", "family dentistry",
    "invisalign", "veneers", "construction", "general contractor", "remodeling",
    "renovation", "roofing", "concrete", "commercial construction",
    "residential construction", "design build", "project management", "home building",
    "excavation", "painting", "plumbing", "electrical"
]


def get_next_output_file(output_dir: Path, base_name: str = "business_leads_scrape", extension: str = ".xlsx") -> Path:
    today = datetime.now().strftime("%m.%d")
    existing_files = list(output_dir.glob(f"{base_name}_{today}_*{extension}"))
    versions = []
    for file in existing_files:
        match = re.search(rf"^{re.escape(base_name)}_{re.escape(today)}_(\d+)$", file.stem)
        if match:
            versions.append(int(match.group(1)))
    next_version = max(versions, default=0) + 1
    return output_dir / f"{base_name}_{today}_{next_version}{extension}"


def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def is_valid_url(url):
    return url.startswith("http://") or url.startswith("https://")


def normalize_url(url):
    url = clean_text(url)
    if not url:
        return ""
    if not url.startswith("http://") and not url.startswith("https://"):
        return "https://" + url
    return url


def domain_of(url):
    try:
        return urlparse(url).netloc.lower().replace("www.", "")
    except Exception:
        return ""


def same_domain(url_a, url_b):
    return domain_of(url_a) == domain_of(url_b)


def fetch_page(url):
    headers = {"User-Agent": "Mozilla/5.0 Business Lead Scraper for internal research"}
    try:
        response = requests.get(url, headers=headers, timeout=15, allow_redirects=True)
        return response, ""
    except requests.RequestException as exc:
        return None, str(exc)


def soup_from_response(response):
    content_type = response.headers.get("Content-Type", "")
    if "text/html" not in content_type:
        return None
    return BeautifulSoup(response.text, "html.parser")


def unique_join(values, limit=10):
    cleaned = []
    seen = set()
    for value in values:
        value = clean_text(value)
        if not value:
            continue
        key = value.lower()
        if key not in seen:
            cleaned.append(value)
            seen.add(key)
    return ", ".join(cleaned[:limit])


def add_issue(issues, issue):
    if issue not in issues:
        issues.append(issue)


def extract_emails(text, soup):
    found = re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
    for link in soup.find_all("a", href=True):
        href = link.get("href", "")
        if href.startswith("mailto:"):
            found.append(href.replace("mailto:", "").split("?")[0])
    return sorted(set(email.lower() for email in found))


def extract_phone_numbers(text, soup):
    found = re.findall(r"(?:\+?\d[\d\s().-]{7,}\d)", text)
    for link in soup.find_all("a", href=True):
        href = link.get("href", "")
        if href.startswith("tel:"):
            found.append(href.replace("tel:", ""))
    cleaned = []
    for phone in found:
        phone = phone.strip()
        digits = re.sub(r"\D", "", phone)
        if 8 <= len(digits) <= 15:
            cleaned.append(phone)
    return sorted(set(cleaned))


def extract_address_candidates(text):
    candidates = []
    pattern = r"\b\d{2,6}\s+[A-Za-z0-9\s.,#'-]+(?:Street|St\.?|Avenue|Ave\.?|Road|Rd\.?|Boulevard|Blvd\.?|Drive|Dr\.?|Lane|Ln\.?|Court|Ct\.?|Way|Highway|Hwy\.?|Suite|Ste\.?)\s*[A-Za-z0-9\s.,#'-]*"
    candidates += re.findall(pattern, text, flags=re.IGNORECASE)
    for line in [line.strip() for line in text.split("\n") if line.strip()]:
        lower = line.lower()
        if any(word in lower for word in ["street", " st ", "avenue", " ave ", "road", " rd ", "suite", "ste.", "blvd", "drive"]):
            if 10 <= len(line) <= 180:
                candidates.append(line)
    return sorted(set(candidates))


def extract_people_candidates(text):
    candidates = []
    patterns = [
        r"(?:CEO|Founder|Owner|Principal|President|Managing Director|Director|Practice Owner|Doctor|Dr\.)\s*[:\-]?\s*([A-Z][A-Za-z'.-]+(?:\s+[A-Z][A-Za-z'.-]+){0,3})",
        r"([A-Z][A-Za-z'.-]+(?:\s+[A-Z][A-Za-z'.-]+){0,3})\s*,?\s*(?:CEO|Founder|Owner|Principal|President|Managing Director|Director|Practice Owner|DDS|DMD)",
        r"(?:Meet|About)\s+(?:Dr\.\s*)?([A-Z][A-Za-z'.-]+(?:\s+[A-Z][A-Za-z'.-]+){0,3})",
    ]
    for pattern in patterns:
        candidates += re.findall(pattern, text)
    blocked = {"Privacy Policy", "Terms Conditions", "Contact Us", "Read More", "Learn More", "Book Now", "Get Started"}
    return sorted(set(name.strip() for name in candidates if name.strip() and name.strip() not in blocked and len(name.split()) <= 4))


def extract_social_links(soup, base_url):
    links = []
    for link in soup.find_all("a", href=True):
        href = urljoin(base_url, link.get("href", "").strip())
        if any(domain in href.lower() for domain in SOCIAL_DOMAINS):
            links.append(href)
    return sorted(set(links))


def extract_google_maps_links(soup, base_url):
    links = []
    for link in soup.find_all("a", href=True):
        href = urljoin(base_url, link.get("href", "").strip())
        lower = href.lower()
        if "google.com/maps" in lower or "goo.gl/maps" in lower:
            links.append(href)
    return sorted(set(links))


def detect_services(text):
    lower_text = text.lower()
    return sorted(set(keyword.title() for keyword in SERVICE_KEYWORDS if keyword in lower_text))


def get_site_name(soup, fallback_url):
    og_site = soup.find("meta", property="og:site_name")
    if og_site and og_site.get("content"):
        return clean_text(og_site.get("content"))
    title = soup.find("title")
    if title:
        title_text = title.get_text(" ", strip=True)
        if title_text:
            return title_text.split("|")[0].split("-")[0].strip()
    return domain_of(fallback_url)


def score_internal_link(href, text):
    combined = f"{href} {text}".lower()
    score = sum(5 for keyword in IMPORTANT_PAGE_KEYWORDS if keyword in combined)
    if any(word in combined for word in ["privacy", "terms", "login", "cart", "checkout", "wp-admin", "tag/", "author/"]):
        score -= 10
    return score


def discover_important_pages(home_url, home_soup):
    discovered = []
    seen = set()
    for link in home_soup.find_all("a", href=True):
        href_raw = link.get("href", "").strip()
        text = link.get_text(" ", strip=True)
        if not href_raw or href_raw.startswith("#") or href_raw.startswith("mailto:") or href_raw.startswith("tel:"):
            continue
        full_url = urljoin(home_url, href_raw).split("#")[0].rstrip("/")
        parsed = urlparse(full_url)
        if parsed.scheme not in ["http", "https"] or not same_domain(home_url, full_url) or full_url in seen:
            continue
        score = score_internal_link(full_url, text)
        if score > 0:
            discovered.append((score, full_url))
            seen.add(full_url)
    discovered = sorted(discovered, key=lambda item: item[0], reverse=True)
    pages = [home_url.rstrip("/")]
    for _, url in discovered:
        if url not in pages:
            pages.append(url)
        if len(pages) >= MAX_PAGES_PER_SITE:
            break
    return pages


def scrape_business_site(start_url, business_type="", notes=""):
    issues = []
    start_url = normalize_url(start_url)
    result = {
        "Website URL": start_url,
        "Final URL": "",
        "Site Name": "",
        "Business Type": business_type,
        "Pages Crawled": 0,
        "Important Pages Crawled": "",
        "Emails": "",
        "Phone Numbers": "",
        "Address Candidates": "",
        "Owner / CEO / Founder Candidates": "",
        "Contact Page": "",
        "About Page": "",
        "Services Keywords": "",
        "Social Links": "",
        "Google Maps Links": "",
        "Forms Detected": "No",
        "Issues": "",
        "Notes": notes,
    }
    if not start_url:
        return None, {"URL": start_url, "Issue": "Missing URL"}
    if not is_valid_url(start_url):
        return None, {"URL": start_url, "Issue": "Invalid URL format"}
    response, error = fetch_page(start_url)
    if response is None:
        return None, {"URL": start_url, "Issue": f"Request failed: {error}"}
    if response.status_code >= 400:
        return None, {"URL": start_url, "Issue": f"HTTP error {response.status_code}"}
    soup = soup_from_response(response)
    if soup is None:
        return None, {"URL": start_url, "Issue": "URL did not return HTML content"}

    final_url = response.url.rstrip("/")
    result["Final URL"] = final_url
    result["Site Name"] = get_site_name(soup, final_url)
    pages_to_crawl = discover_important_pages(final_url, soup)

    all_emails, all_phones, all_addresses, all_people = [], [], [], []
    all_socials, all_maps, all_services, crawled_pages = [], [], [], []
    contact_page, about_page = "", ""
    form_detected = False

    for page_url in pages_to_crawl:
        time.sleep(REQUEST_DELAY_SECONDS)
        page_response, _ = fetch_page(page_url)
        if page_response is None or page_response.status_code >= 400:
            continue
        page_soup = soup_from_response(page_response)
        if page_soup is None:
            continue
        page_url_final = page_response.url.rstrip("/")
        crawled_pages.append(page_url_final)
        lower_page_url = page_url_final.lower()
        if not contact_page and "contact" in lower_page_url:
            contact_page = page_url_final
        if not about_page and any(word in lower_page_url for word in ["about", "team", "leadership", "staff", "doctor", "dentist"]):
            about_page = page_url_final
        text = page_soup.get_text("\n", strip=True)
        all_emails += extract_emails(text, page_soup)
        all_phones += extract_phone_numbers(text, page_soup)
        all_addresses += extract_address_candidates(text)
        all_people += extract_people_candidates(text)
        all_socials += extract_social_links(page_soup, page_url_final)
        all_maps += extract_google_maps_links(page_soup, page_url_final)
        all_services += detect_services(text)
        if page_soup.find_all("form"):
            form_detected = True

    result["Pages Crawled"] = len(crawled_pages)
    result["Important Pages Crawled"] = unique_join(crawled_pages, limit=20)
    result["Emails"] = unique_join(all_emails, limit=10)
    result["Phone Numbers"] = unique_join(all_phones, limit=10)
    result["Address Candidates"] = unique_join(all_addresses, limit=5)
    result["Owner / CEO / Founder Candidates"] = unique_join(all_people, limit=8)
    result["Contact Page"] = contact_page
    result["About Page"] = about_page
    result["Services Keywords"] = unique_join(all_services, limit=15)
    result["Social Links"] = unique_join(all_socials, limit=10)
    result["Google Maps Links"] = unique_join(all_maps, limit=5)
    result["Forms Detected"] = "Yes" if form_detected else "No"
    if not result["Emails"]: add_issue(issues, "No email found")
    if not result["Phone Numbers"]: add_issue(issues, "No phone number found")
    if not result["Address Candidates"]: add_issue(issues, "No address candidate found")
    if not result["Owner / CEO / Founder Candidates"]: add_issue(issues, "No owner/CEO/founder candidate found")
    if not result["Contact Page"]: add_issue(issues, "No contact page detected")
    if len(crawled_pages) <= 1: add_issue(issues, "Only homepage crawled")
    result["Issues"] = "; ".join(issues)
    return result, None


def build_summary(leads_df, invalid_df):
    return pd.DataFrame([
        ["Input file", str(INPUT_FILE)],
        ["Businesses processed", len(leads_df) + len(invalid_df)],
        ["Valid websites scraped", len(leads_df)],
        ["Invalid URLs", len(invalid_df)],
        ["Sites with email found", int(leads_df["Emails"].astype(str).str.strip().ne("").sum()) if len(leads_df) else 0],
        ["Sites with phone found", int(leads_df["Phone Numbers"].astype(str).str.strip().ne("").sum()) if len(leads_df) else 0],
        ["Sites with address candidate found", int(leads_df["Address Candidates"].astype(str).str.strip().ne("").sum()) if len(leads_df) else 0],
        ["Sites with owner/CEO/founder candidate found", int(leads_df["Owner / CEO / Founder Candidates"].astype(str).str.strip().ne("").sum()) if len(leads_df) else 0],
        ["Max pages per site", MAX_PAGES_PER_SITE],
        ["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ], columns=["Metric", "Value"])


def get_col(ws, name):
    for cell in ws[1]:
        if cell.value == name:
            return cell.column
    return None


def format_workbook(output_file):
    wb = load_workbook(output_file)
    dark_fill = PatternFill("solid", fgColor="111827")
    green_fill = PatternFill("solid", fgColor="166534")
    red_fill = PatternFill("solid", fgColor="991B1B")
    summary_fill = PatternFill("solid", fgColor="E0F2FE")
    issue_fill = PatternFill("solid", fgColor="F4CCCC")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    white_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    for ws in wb.worksheets:
        max_row, max_col = ws.max_row, ws.max_column
        if ws.title == "Summary":
            ws.freeze_panes = None
        else:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = dark_fill
            cell.font = white_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if ws.title == "Business_Leads":
            for cell in ws[1]: cell.fill = green_fill
        if ws.title == "Invalid_URLs":
            for cell in ws[1]: cell.fill = red_fill
        if ws.title == "Summary":
            for row in range(2, max_row + 1): ws[f"A{row}"].fill = summary_fill
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        issues_col = get_col(ws, "Issues")
        if issues_col:
            ws.cell(1, issues_col).fill = PatternFill("solid", fgColor="D99694")
            ws.cell(1, issues_col).font = Font(color="000000", bold=True)
            for row in range(2, max_row + 1):
                issues = str(ws.cell(row, issues_col).value or "")
                if issues:
                    ws.cell(row, issues_col).fill = issue_fill
                    for cell in ws[row]:
                        if cell.fill.fill_type is None:
                            cell.fill = warning_fill
                    ws.cell(row, issues_col).fill = issue_fill
        if max_row >= 2 and max_col >= 1 and ws.title != "Summary":
            table_name = re.sub(r"[^A-Za-z0-9_]", "", ws.title)[:20] + "Table"
            ref = f"A1:{get_column_letter(max_col)}{max_row}"
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            try: ws.add_table(table)
            except ValueError: pass
        for col_idx in range(1, max_col + 1):
            header = str(ws.cell(1, col_idx).value or "")
            width = 18
            if header in ["Website URL", "Final URL", "Important Pages Crawled", "Social Links", "Google Maps Links"]: width = 55
            elif header in ["Emails", "Phone Numbers", "Address Candidates", "Owner / CEO / Founder Candidates", "Services Keywords", "Issues", "Notes"]: width = 45
            elif header in ["Site Name", "Business Type", "Contact Page", "About Page"]: width = 28
            elif header == "URL": width = 55
            elif header == "Issue": width = 45
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        for row in range(1, max_row + 1): ws.row_dimensions[row].height = 26
    wb.save(output_file)


def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Input file not found:\n{INPUT_FILE}\n\nPlace business_websites.xlsx inside the datasets folder.")
    output_file = get_next_output_file(OUTPUTS_DIR)
    print("Reading business website list...")
    df = pd.read_excel(INPUT_FILE, sheet_name="Websites_Raw")
    df.columns = [str(col).strip() for col in df.columns]
    if "URL" not in df.columns:
        raise ValueError("Missing required column: URL")
    if "Business Type" not in df.columns: df["Business Type"] = ""
    if "Notes" not in df.columns: df["Notes"] = ""
    lead_rows, invalid_rows = [], []
    print("Scraping business websites...")
    for _, row in df.iterrows():
        url = clean_text(row["URL"])
        print(f"Scraping: {url}")
        lead, invalid = scrape_business_site(url, clean_text(row["Business Type"]), clean_text(row["Notes"]))
        if invalid: invalid_rows.append(invalid)
        else: lead_rows.append(lead)
    leads_df = pd.DataFrame(lead_rows)
    invalid_df = pd.DataFrame(invalid_rows, columns=["URL", "Issue"])
    summary_df = build_summary(leads_df, invalid_df)
    print("Writing business lead scraper report...")
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        leads_df.to_excel(writer, sheet_name="Business_Leads", index=False)
        invalid_df.to_excel(writer, sheet_name="Invalid_URLs", index=False)
    print("Applying Excel formatting...")
    format_workbook(output_file)
    print("\nBusiness website scraping complete.")
    print(f"Created: {output_file.name}")
    print(f"Saved in: {OUTPUTS_DIR}")
    print(f"Valid websites scraped: {len(leads_df)}")
    print(f"Invalid URLs: {len(invalid_df)}")


if __name__ == "__main__":
    main()
