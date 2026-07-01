"""
clean_validate_excel_formatted.py

Project:
Clean and Validate Excel / CSV

Purpose:
Clean customer and order data from large_messy_dataset.xlsx, split valid and invalid rows,
and generate a formatted Excel report with highlighted issue cells.

Expected folder structure:

Python/
    datasets/
        large_messy_dataset.xlsx
    scripts/
        clean_validate_excel_formatted.py
    outputs/
    archive/

Run from the Python folder:
    py scripts/clean_validate_excel_formatted.py

Output format:
    outputs/cleaned_dataset_06.30_1.xlsx
    outputs/cleaned_dataset_06.30_2.xlsx

Output sheets:
    Summary
    Customers_Valid
    Customers_Invalid
    Orders_Valid
    Orders_Invalid

Main features:
    - Cleans names, emails, phone numbers, dates, countries, statuses
    - Cleans product names, quantities, prices, dates, payment statuses
    - Moves invalid rows into _Invalid sheets
    - Adds final Issues column to invalid sheets
    - Highlights issue-related cells
    - Formats workbook for portfolio/client presentation
"""

import re
from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent.parent
DATASETS_DIR = BASE_DIR / "datasets"
OUTPUTS_DIR = BASE_DIR / "outputs"
ARCHIVE_DIR = BASE_DIR / "archive"

INPUT_FILE = DATASETS_DIR / "large_messy_dataset.xlsx"

DATASETS_DIR.mkdir(exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)
ARCHIVE_DIR.mkdir(exist_ok=True)


# ---------- File naming ----------

def get_next_output_file(output_dir: Path, base_name: str = "cleaned_dataset", extension: str = ".xlsx") -> Path:
    today = datetime.now().strftime("%m.%d")
    existing_files = list(output_dir.glob(f"{base_name}_{today}_*{extension}"))

    versions = []
    for file in existing_files:
        match = re.search(rf"^{re.escape(base_name)}_{re.escape(today)}_(\d+)$", file.stem)
        if match:
            versions.append(int(match.group(1)))

    next_version = max(versions, default=0) + 1
    return output_dir / f"{base_name}_{today}_{next_version}{extension}"


# ---------- Cleaning helpers ----------

def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def normalize_name(value):
    return clean_text(value).title()


def normalize_email(value):
    return clean_text(value).lower()


def is_valid_email(value):
    if not value:
        return False
    return bool(re.match(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$", value))


def normalize_phone(value):
    value = clean_text(value)
    digits = re.sub(r"\D", "", value)

    if not digits:
        return ""

    if digits.startswith("63") and len(digits) == 12:
        return "+" + digits

    if digits.startswith("09") and len(digits) == 11:
        return "+63" + digits[1:]

    if digits.startswith("9") and len(digits) == 10:
        return "+63" + digits

    return value


def normalize_country(value):
    value = clean_text(value).lower()
    if value in ["ph", "philippines", "philippine"]:
        return "Philippines"
    return value.title() if value else ""


def normalize_status(value):
    value = clean_text(value).lower()
    mapping = {
        "active": "Active",
        "inactive": "Inactive",
        "pending": "Pending",
        "paid": "Paid",
        "refunded": "Refunded",
        "failed": "Failed",
    }
    return mapping.get(value, value.title() if value else "")


def parse_date(value):
    if pd.isna(value) or str(value).strip() == "":
        return ""

    parsed = pd.to_datetime(value, errors="coerce", dayfirst=False)

    if pd.isna(parsed):
        parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)

    if pd.isna(parsed):
        return ""

    return parsed.strftime("%Y-%m-%d")


def clean_price(value):
    if pd.isna(value) or str(value).strip() == "":
        return None

    cleaned = str(value).replace(",", "").strip()

    try:
        return float(cleaned)
    except ValueError:
        return None


def clean_quantity(value):
    if pd.isna(value) or str(value).strip() == "":
        return None

    try:
        return int(value)
    except ValueError:
        return None


def add_issue(issues_by_index, row_index, issue):
    if row_index not in issues_by_index:
        issues_by_index[row_index] = []
    issues_by_index[row_index].append(issue)


def count_total_issues(issues_by_index):
    return sum(len(issues) for issues in issues_by_index.values())


def create_duplicate_group_ids(df, duplicate_columns):
    duplicate_group_ids = pd.Series("", index=df.index, dtype="object")

    valid_key_mask = pd.Series(True, index=df.index)
    for col in duplicate_columns:
        valid_key_mask = valid_key_mask & df[col].astype(str).str.strip().ne("")

    duplicated_mask = df[valid_key_mask].duplicated(subset=duplicate_columns, keep=False)
    duplicate_rows = df[valid_key_mask].loc[duplicated_mask].copy()

    if duplicate_rows.empty:
        return duplicate_group_ids

    group_number = 1
    for _, group in duplicate_rows.groupby(duplicate_columns, dropna=False, sort=True):
        for index in group.index:
            duplicate_group_ids.loc[index] = f"DUP-{group_number}"
        group_number += 1

    return duplicate_group_ids


def split_valid_invalid(df, issues_by_index, sort_columns=None, duplicate_group_ids=None):
    invalid_indexes = sorted(issues_by_index.keys())

    valid_df = df.drop(index=invalid_indexes, errors="ignore").copy()
    invalid_df = df.loc[df.index.intersection(invalid_indexes)].copy()

    if len(invalid_df) > 0:
        invalid_df["Issues"] = invalid_df.index.map(
            lambda idx: "; ".join(sorted(set(issues_by_index.get(idx, []))))
        )

        if duplicate_group_ids is not None:
            invalid_df["_duplicate_sort_key"] = invalid_df.index.map(duplicate_group_ids).fillna("")
        else:
            invalid_df["_duplicate_sort_key"] = ""

        invalid_df["_original_index"] = invalid_df.index

        available_sort_columns = [
            col for col in (sort_columns or [])
            if col in invalid_df.columns
        ]

        invalid_df = invalid_df.sort_values(
            by=available_sort_columns + ["_duplicate_sort_key", "_original_index"],
            kind="stable"
        ).drop(columns=["_duplicate_sort_key", "_original_index"])
    else:
        invalid_df["Issues"] = []

    return valid_df.reset_index(drop=True), invalid_df.reset_index(drop=True)


# ---------- Customer cleaning ----------

def clean_customers(customers):
    issues_by_index = {}

    customers = customers.copy()
    customers.columns = [str(col).strip() for col in customers.columns]

    customers["Customer ID"] = customers["Customer ID"].apply(clean_text)
    customers["Name"] = customers["Name"].apply(normalize_name)
    customers["Email"] = customers["Email"].apply(normalize_email)
    customers["Phone"] = customers["Phone"].apply(normalize_phone)
    customers["City"] = customers["City"].apply(lambda x: clean_text(x).title())
    customers["Country"] = customers["Country"].apply(normalize_country)
    customers["Signup Date"] = customers["Signup Date"].apply(parse_date)
    customers["Status"] = customers["Status"].apply(normalize_status)

    for index, row in customers.iterrows():
        if not row["Customer ID"]:
            add_issue(issues_by_index, index, "Missing customer ID")

        if not row["Email"]:
            add_issue(issues_by_index, index, "Missing email")
        elif not is_valid_email(row["Email"]):
            add_issue(issues_by_index, index, "Invalid email format")

        if not row["Phone"]:
            add_issue(issues_by_index, index, "Missing phone")

        if not row["Signup Date"]:
            add_issue(issues_by_index, index, "Missing or invalid signup date")

    # All duplicate Customer ID rows are invalid.
    customer_duplicate_group_ids = create_duplicate_group_ids(customers, ["Customer ID"])
    duplicate_indexes = customer_duplicate_group_ids[customer_duplicate_group_ids.ne("")].index
    for index in duplicate_indexes:
        add_issue(issues_by_index, index, "Duplicate customer ID")

    customers_valid, customers_invalid = split_valid_invalid(
        customers,
        issues_by_index,
        sort_columns=["Customer ID", "Name", "Email"],
        duplicate_group_ids=customer_duplicate_group_ids
    )

    return customers_valid, customers_invalid, issues_by_index


# ---------- Order cleaning ----------

def clean_orders(orders):
    issues_by_index = {}

    orders = orders.copy()
    orders.columns = [str(col).strip() for col in orders.columns]

    orders["Order ID"] = orders["Order ID"].apply(clean_text)
    orders["Customer ID"] = orders["Customer ID"].apply(clean_text)
    orders["Product"] = orders["Product"].apply(lambda x: clean_text(x).title())

    if "Category" in orders.columns:
        orders["Category"] = orders["Category"].apply(lambda x: clean_text(x).title())

    orders["Quantity"] = orders["Quantity"].apply(clean_quantity)
    orders["Price"] = orders["Price"].apply(clean_price)
    orders["Order Date"] = orders["Order Date"].apply(parse_date)
    orders["Payment Status"] = orders["Payment Status"].apply(normalize_status)

    if "Channel" in orders.columns:
        orders["Channel"] = orders["Channel"].apply(clean_text)

    for index, row in orders.iterrows():
        if not row["Order ID"]:
            add_issue(issues_by_index, index, "Missing order ID")

        if not row["Customer ID"]:
            add_issue(issues_by_index, index, "Missing customer ID")

        if pd.isna(row["Quantity"]):
            add_issue(issues_by_index, index, "Missing or invalid quantity")
        elif row["Quantity"] <= 0:
            add_issue(issues_by_index, index, "Quantity is zero or negative")

        if pd.isna(row["Price"]):
            add_issue(issues_by_index, index, "Missing or invalid price")

        if not row["Order Date"]:
            add_issue(issues_by_index, index, "Missing or invalid order date")

    # Order ID alone is valid.
    # Duplicate only if Order ID + Customer ID + Product match.
    duplicate_columns = ["Order ID", "Customer ID", "Product"]
    order_duplicate_group_ids = create_duplicate_group_ids(orders, duplicate_columns)
    duplicate_indexes = order_duplicate_group_ids[order_duplicate_group_ids.ne("")].index
    for index in duplicate_indexes:
        add_issue(issues_by_index, index, "Duplicate order row based on Order ID + Customer ID + Product")

    orders_valid, orders_invalid = split_valid_invalid(
        orders,
        issues_by_index,
        sort_columns=["Customer ID", "Order ID", "Product"],
        duplicate_group_ids=order_duplicate_group_ids
    )

    return orders_valid, orders_invalid, issues_by_index


# ---------- Excel formatting ----------

def get_column_index(ws, column_name):
    for cell in ws[1]:
        if cell.value == column_name:
            return cell.column
    return None


def format_workbook(output_file):
    wb = load_workbook(output_file)

    dark_header_fill = PatternFill("solid", fgColor="1F2937")
    dark_header_font = Font(color="FFFFFF", bold=True)
    summary_fill = PatternFill("solid", fgColor="E0F2FE")
    issue_header_fill = PatternFill("solid", fgColor="D99694")
    issue_fill = PatternFill("solid", fgColor="F4CCCC")
    invalid_cell_fill = PatternFill("solid", fgColor="FCE4E4")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    duplicate_fill = PatternFill("solid", fgColor="FCE5CD")
    valid_header_fill = PatternFill("solid", fgColor="166534")
    invalid_header_fill = PatternFill("solid", fgColor="991B1B")

    thin_gray = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    issue_to_columns = {
        "Missing customer ID": ["Customer ID"],
        "Duplicate customer ID": ["Customer ID"],
        "Missing email": ["Email"],
        "Invalid email format": ["Email"],
        "Missing phone": ["Phone"],
        "Missing or invalid signup date": ["Signup Date"],

        "Missing order ID": ["Order ID"],
        "Missing customer ID": ["Customer ID"],
        "Missing or invalid quantity": ["Quantity"],
        "Quantity is zero or negative": ["Quantity"],
        "Missing or invalid price": ["Price"],
        "Missing or invalid order date": ["Order Date"],
        "Duplicate order row based on Order ID + Customer ID + Product": ["Order ID", "Customer ID", "Product"],
    }

    for ws in wb.worksheets:
        max_row = ws.max_row
        max_col = ws.max_column

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Header formatting
        for cell in ws[1]:
            cell.fill = dark_header_fill
            cell.font = dark_header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        if "Valid" in ws.title:
            for cell in ws[1]:
                cell.fill = valid_header_fill

        if "Invalid" in ws.title:
            for cell in ws[1]:
                cell.fill = invalid_header_fill

        if ws.title == "Summary":
            ws.freeze_panes = None
            for row in ws.iter_rows():
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

            ws["A1"].fill = dark_header_fill
            ws["A1"].font = dark_header_font
            ws["B1"].fill = dark_header_fill
            ws["B1"].font = dark_header_font

            for row in range(2, max_row + 1):
                ws[f"A{row}"].fill = summary_fill

            ws.column_dimensions["A"].width = 34
            ws.column_dimensions["B"].width = 80
            continue

        # General cell formatting
        for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        # Number formats
        for col_name in ["Price"]:
            col_idx = get_column_index(ws, col_name)
            if col_idx:
                for row in range(2, max_row + 1):
                    ws.cell(row=row, column=col_idx).number_format = "#,##0.00"

        for col_name in ["Quantity", "Stock"]:
            col_idx = get_column_index(ws, col_name)
            if col_idx:
                for row in range(2, max_row + 1):
                    ws.cell(row=row, column=col_idx).number_format = "0"

        for col_name in ["Signup Date", "Order Date"]:
            col_idx = get_column_index(ws, col_name)
            if col_idx:
                for row in range(2, max_row + 1):
                    ws.cell(row=row, column=col_idx).number_format = "yyyy-mm-dd"

        # Invalid sheet issue highlighting
        if "Invalid" in ws.title:
            issues_col = get_column_index(ws, "Issues")

            if issues_col:
                ws.cell(row=1, column=issues_col).fill = issue_header_fill
                ws.cell(row=1, column=issues_col).font = Font(color="000000", bold=True)

                for row in range(2, max_row + 1):
                    issues_text = str(ws.cell(row=row, column=issues_col).value or "")
                    ws.cell(row=row, column=issues_col).fill = issue_fill

                    if "Duplicate" in issues_text:
                        for cell in ws[row]:
                            cell.fill = duplicate_fill
                        ws.cell(row=row, column=issues_col).fill = issue_fill

                    if "Low" in issues_text or "Pending" in issues_text:
                        for cell in ws[row]:
                            cell.fill = warning_fill
                        ws.cell(row=row, column=issues_col).fill = issue_fill

                    for issue, columns in issue_to_columns.items():
                        if issue in issues_text:
                            for col_name in columns:
                                col_idx = get_column_index(ws, col_name)
                                if col_idx:
                                    ws.cell(row=row, column=col_idx).fill = invalid_cell_fill

        # Create Excel table
        if max_row >= 2 and max_col >= 1:
            table_name = re.sub(r"[^A-Za-z0-9_]", "", ws.title)[:20] + "Table"
            ref = f"A1:{get_column_letter(max_col)}{max_row}"
            table = Table(displayName=table_name, ref=ref)
            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            try:
                ws.add_table(table)
            except ValueError:
                pass

        # Column widths
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            header = str(ws.cell(row=1, column=col_idx).value or "")

            if header in ["Issues"]:
                ws.column_dimensions[col_letter].width = 55
            elif header in ["Email", "Image URL", "Product URL", "Description"]:
                ws.column_dimensions[col_letter].width = 36
            elif header in ["Name", "Product"]:
                ws.column_dimensions[col_letter].width = 24
            elif header in ["Customer ID", "Order ID"]:
                ws.column_dimensions[col_letter].width = 18
            else:
                ws.column_dimensions[col_letter].width = 16

        for row in range(1, max_row + 1):
            ws.row_dimensions[row].height = 22

    wb.save(output_file)


# ---------- Main ----------

def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Input file not found:\n{INPUT_FILE}\n\n"
            "Place large_messy_dataset.xlsx inside the datasets folder."
        )

    output_file = get_next_output_file(OUTPUTS_DIR)

    print("Reading workbook...")
    print(f"Input: {INPUT_FILE}")

    customers = pd.read_excel(INPUT_FILE, sheet_name="Customers_Raw")
    orders = pd.read_excel(INPUT_FILE, sheet_name="Orders_Raw")

    print("Cleaning customers...")
    customers_valid, customers_invalid, customer_issues = clean_customers(customers)

    print("Cleaning orders...")
    orders_valid, orders_invalid, order_issues = clean_orders(orders)

    summary = pd.DataFrame([
        ["Input file", str(INPUT_FILE)],
        ["Output file", str(output_file)],
        ["Raw customer rows", len(customers)],
        ["Customers valid rows", len(customers_valid)],
        ["Customers invalid rows moved", len(customers_invalid)],
        ["Customer issue count", count_total_issues(customer_issues)],
        ["Raw order rows", len(orders)],
        ["Orders valid rows", len(orders_valid)],
        ["Orders invalid rows moved", len(orders_invalid)],
        ["Order issue count", count_total_issues(order_issues)],
        ["Customer duplicate rule", "All rows are invalid when Customer ID appears more than once"],
        ["Order duplicate rule", "Duplicate only when Order ID + Customer ID + Product match"],
        ["Invalid sorting", "Invalid sheets are sorted so duplicates appear consecutively"],
        ["Formatting", "Issue cells and Issues column are highlighted automatically"],
        ["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ], columns=["Metric", "Value"])

    print("Writing output workbook...")
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        customers_valid.to_excel(writer, sheet_name="Customers_Valid", index=False)
        customers_invalid.to_excel(writer, sheet_name="Customers_Invalid", index=False)
        orders_valid.to_excel(writer, sheet_name="Orders_Valid", index=False)
        orders_invalid.to_excel(writer, sheet_name="Orders_Invalid", index=False)

    print("Applying Excel formatting...")
    format_workbook(output_file)

    print("\nCleaning complete.")
    print(f"Created: {output_file.name}")
    print(f"Saved in: {OUTPUTS_DIR}")
    print(f"Customers valid rows: {len(customers_valid)}")
    print(f"Customers invalid rows moved: {len(customers_invalid)}")
    print(f"Customer issue count: {count_total_issues(customer_issues)}")
    print(f"Orders valid rows: {len(orders_valid)}")
    print(f"Orders invalid rows moved: {len(orders_invalid)}")
    print(f"Order issue count: {count_total_issues(order_issues)}")


if __name__ == "__main__":
    main()
