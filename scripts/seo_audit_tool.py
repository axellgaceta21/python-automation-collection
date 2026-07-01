"""
seo_audit_tool.py

Project:
SEO Audit Tool

Purpose:
Read a list of URLs, check common technical/on-page SEO elements,
and generate a formatted Excel audit report.

Expected folder structure:

Python/
    datasets/
        seo_urls_sample.xlsx
    scripts/
        seo_audit_tool.py
    outputs/
    archive/

Run from the Python folder:
    py scripts/seo_audit_tool.py

Required packages:
    py -m pip install pandas openpyxl requests beautifulsoup4

Output format:
    outputs/seo_audit_06.30_1.xlsx
    outputs/seo_audit_06.30_2.xlsx

Output sheets:
    Summary
    SEO_Audit
    URLs_Invalid

Main checks:
    - URL format
    - HTTP status code
    - HTTPS usage
    - Title tag presence and length
    - Meta description presence and length
    - H1 count
    - Canonical tag
    - Robots meta tag
    - Image alt text counts
    - Word count estimate
    - Final Issues column
"""

import re
from pathlib import Path
from datetime import datetime

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent.parent
DATASETS_DIR = BASE_DIR / "datasets"
OUTPUTS_DIR = BASE_DIR / "outputs"
ARCHIVE_DIR = BASE_DIR / "archive"

INPUT_FILE = DATASETS_DIR / "seo_urls_sample.xlsx"

DATASETS_DIR.mkdir(exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)
ARCHIVE_DIR.mkdir(exist_ok=True)


def get_next_output_file(output_dir: Path, base_name: str = "seo_audit", extension: str = ".xlsx") -> Path:
    today = datetime.now().strftime("%m.%d")
    existing_files = list(output_dir.glob(f"{base_name}_{today}_*{extension}"))

    versions = []
    for file in existing_files:
        match = re.search(rf"^{re.escape(base_name)}_{re.escape(today)}_(\d+)$", file.stem)
        if match:
            versions.append(int(match.group(1)))

    next_version = max(versions, default=0) + 1
    return output_dir / f"{base_name}_{today}_{next_version}{extension}"


def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def is_valid_url(url):
    return url.startswith("http://") or url.startswith("https://")


def add_issue(issues, issue):
    if issue not in issues:
        issues.append(issue)


def fetch_url(url):
    headers = {
        "User-Agent": "Mozilla/5.0 SEO Audit Bot for internal technical QA"
    }

    try:
        response = requests.get(url, headers=headers, timeout=12, allow_redirects=True)
        return response, ""
    except requests.RequestException as exc:
        return None, str(exc)


def audit_single_url(url, page_type="", priority="", notes=""):
    issues = []

    result = {
        "URL": url,
        "Final URL": "",
        "Page Type": page_type,
        "Priority": priority,
        "HTTP Status": "",
        "HTTPS": "",
        "Title": "",
        "Title Length": "",
        "Meta Description": "",
        "Meta Description Length": "",
        "H1 Count": "",
        "H1 Text": "",
        "Canonical": "",
        "Robots Meta": "",
        "Images": "",
        "Images Missing Alt": "",
        "Word Count": "",
        "Issues": "",
        "Notes": notes,
    }

    if not url:
        result["Issues"] = "Missing URL"
        return result

    if not is_valid_url(url):
        result["Issues"] = "Invalid URL format"
        return result

    if url.startswith("http://"):
        add_issue(issues, "URL is not HTTPS")

    response, error = fetch_url(url)

    if response is None:
        result["Issues"] = f"Request failed: {error}"
        return result

    result["Final URL"] = response.url
    result["HTTP Status"] = response.status_code
    result["HTTPS"] = "Yes" if response.url.startswith("https://") else "No"

    if response.status_code >= 400:
        add_issue(issues, f"HTTP error {response.status_code}")

    content_type = response.headers.get("Content-Type", "")
    if "text/html" not in content_type:
        add_issue(issues, "URL did not return HTML content")
        result["Issues"] = "; ".join(issues)
        return result

    soup = BeautifulSoup(response.text, "html.parser")

    title_tag = soup.find("title")
    title = title_tag.get_text(" ", strip=True) if title_tag else ""
    result["Title"] = title
    result["Title Length"] = len(title)

    if not title:
        add_issue(issues, "Missing title tag")
    elif len(title) < 30:
        add_issue(issues, "Title too short")
    elif len(title) > 65:
        add_issue(issues, "Title too long")

    meta_desc_tag = soup.find("meta", attrs={"name": re.compile("^description$", re.I)})
    meta_description = meta_desc_tag.get("content", "").strip() if meta_desc_tag else ""
    result["Meta Description"] = meta_description
    result["Meta Description Length"] = len(meta_description)

    if not meta_description:
        add_issue(issues, "Missing meta description")
    elif len(meta_description) < 70:
        add_issue(issues, "Meta description too short")
    elif len(meta_description) > 160:
        add_issue(issues, "Meta description too long")

    h1_tags = soup.find_all("h1")
    h1_text = " | ".join([h.get_text(" ", strip=True) for h in h1_tags[:3]])
    result["H1 Count"] = len(h1_tags)
    result["H1 Text"] = h1_text

    if len(h1_tags) == 0:
        add_issue(issues, "Missing H1")
    elif len(h1_tags) > 1:
        add_issue(issues, "Multiple H1 tags")

    canonical_tag = soup.find("link", attrs={"rel": lambda value: value and "canonical" in value})
    canonical = canonical_tag.get("href", "").strip() if canonical_tag else ""
    result["Canonical"] = canonical

    if not canonical:
        add_issue(issues, "Missing canonical tag")

    robots_tag = soup.find("meta", attrs={"name": re.compile("^robots$", re.I)})
    robots = robots_tag.get("content", "").strip() if robots_tag else ""
    result["Robots Meta"] = robots

    if "noindex" in robots.lower():
        add_issue(issues, "Page has noindex robots meta")

    images = soup.find_all("img")
    missing_alt = 0
    for img in images:
        alt = img.get("alt")
        if alt is None or str(alt).strip() == "":
            missing_alt += 1

    result["Images"] = len(images)
    result["Images Missing Alt"] = missing_alt

    if missing_alt > 0:
        add_issue(issues, "Images missing alt text")

    body_text = soup.get_text(" ", strip=True)
    words = re.findall(r"\w+", body_text)
    result["Word Count"] = len(words)

    if len(words) < 250:
        add_issue(issues, "Low word count")

    result["Issues"] = "; ".join(issues)
    return result


def build_summary(audit_df, invalid_df):
    total_urls = len(audit_df) + len(invalid_df)
    audited_urls = len(audit_df)
    issue_rows = audit_df["Issues"].astype(str).str.strip().ne("").sum()
    invalid_urls = len(invalid_df)
    avg_title_length = audit_df["Title Length"].replace("", pd.NA).dropna().astype(float).mean() if audited_urls else 0
    avg_meta_length = audit_df["Meta Description Length"].replace("", pd.NA).dropna().astype(float).mean() if audited_urls else 0
    missing_h1 = audit_df["Issues"].astype(str).str.contains("Missing H1", na=False).sum() if audited_urls else 0
    missing_meta = audit_df["Issues"].astype(str).str.contains("Missing meta description", na=False).sum() if audited_urls else 0
    missing_canonical = audit_df["Issues"].astype(str).str.contains("Missing canonical tag", na=False).sum() if audited_urls else 0

    return pd.DataFrame([
        ["Input file", str(INPUT_FILE)],
        ["Total input rows", total_urls],
        ["Audited URLs", audited_urls],
        ["Invalid URL rows", invalid_urls],
        ["Audited URLs with issues", int(issue_rows)],
        ["Missing meta descriptions", int(missing_meta)],
        ["Missing H1", int(missing_h1)],
        ["Missing canonical tag", int(missing_canonical)],
        ["Average title length", round(avg_title_length, 1) if pd.notna(avg_title_length) else 0],
        ["Average meta description length", round(avg_meta_length, 1) if pd.notna(avg_meta_length) else 0],
        ["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ], columns=["Metric", "Value"])


def get_col(ws, name):
    for cell in ws[1]:
        if cell.value == name:
            return cell.column
    return None


def format_workbook(output_file):
    wb = load_workbook(output_file)

    dark_fill = PatternFill("solid", fgColor="111827")
    blue_fill = PatternFill("solid", fgColor="1D4ED8")
    red_fill = PatternFill("solid", fgColor="991B1B")
    green_fill = PatternFill("solid", fgColor="166534")
    issue_fill = PatternFill("solid", fgColor="F4CCCC")
    invalid_cell_fill = PatternFill("solid", fgColor="FCE4E4")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    summary_fill = PatternFill("solid", fgColor="E0F2FE")

    white_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    issue_map = {
        "Missing URL": ["URL"],
        "Invalid URL format": ["URL"],
        "Request failed": ["URL"],
        "HTTP error": ["HTTP Status"],
        "URL is not HTTPS": ["URL", "HTTPS"],
        "Missing title tag": ["Title"],
        "Title too short": ["Title", "Title Length"],
        "Title too long": ["Title", "Title Length"],
        "Missing meta description": ["Meta Description"],
        "Meta description too short": ["Meta Description", "Meta Description Length"],
        "Meta description too long": ["Meta Description", "Meta Description Length"],
        "Missing H1": ["H1 Count", "H1 Text"],
        "Multiple H1 tags": ["H1 Count", "H1 Text"],
        "Missing canonical tag": ["Canonical"],
        "Page has noindex robots meta": ["Robots Meta"],
        "Images missing alt text": ["Images Missing Alt"],
        "Low word count": ["Word Count"],
    }

    for ws in wb.worksheets:
        max_row = ws.max_row
        max_col = ws.max_column

        if ws.title == "Summary":
            ws.freeze_panes = None
        else:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = dark_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        if ws.title == "SEO_Audit":
            for cell in ws[1]:
                cell.fill = green_fill

        if ws.title == "URLs_Invalid":
            for cell in ws[1]:
                cell.fill = red_fill

        if ws.title == "Summary":
            for row in range(2, max_row + 1):
                ws[f"A{row}"].fill = summary_fill

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        issues_col = get_col(ws, "Issues")
        if issues_col:
            ws.cell(1, issues_col).fill = PatternFill("solid", fgColor="D99694")
            ws.cell(1, issues_col).font = Font(color="000000", bold=True)

            for row in range(2, max_row + 1):
                issues = str(ws.cell(row, issues_col).value or "")
                if issues:
                    ws.cell(row, issues_col).fill = issue_fill

                    for key, cols in issue_map.items():
                        if key in issues:
                            for col_name in cols:
                                col = get_col(ws, col_name)
                                if col:
                                    ws.cell(row, col).fill = invalid_cell_fill

                    if "too short" in issues or "Low word count" in issues:
                        for cell in ws[row]:
                            if cell.fill.fill_type is None:
                                cell.fill = warning_fill
                        ws.cell(row, issues_col).fill = issue_fill

        for col_name in ["HTTP Status", "Title Length", "Meta Description Length", "H1 Count", "Images", "Images Missing Alt", "Word Count"]:
            col = get_col(ws, col_name)
            if col:
                for row in range(2, max_row + 1):
                    ws.cell(row, col).number_format = "#,##0"

        if max_row >= 2 and max_col >= 1 and ws.title != "Summary":
            table_name = re.sub(r"[^A-Za-z0-9_]", "", ws.title)[:20] + "Table"
            ref = f"A1:{get_column_letter(max_col)}{max_row}"
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            try:
                ws.add_table(table)
            except ValueError:
                pass

        for col_idx in range(1, max_col + 1):
            header = str(ws.cell(1, col_idx).value or "")
            width = 16
            if header in ["URL", "Final URL", "Canonical"]:
                width = 54
            elif header in ["Title", "Meta Description", "H1 Text", "Issues", "Notes"]:
                width = 44
            elif header in ["Page Type", "Priority", "HTTPS", "HTTP Status"]:
                width = 16
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        for row in range(1, max_row + 1):
            ws.row_dimensions[row].height = 22

    wb.save(output_file)


def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Input file not found:\n{INPUT_FILE}\n\n"
            "Place seo_urls_sample.xlsx inside the datasets folder."
        )

    output_file = get_next_output_file(OUTPUTS_DIR)

    print("Reading URL list...")
    df = pd.read_excel(INPUT_FILE, sheet_name="URLs_Raw")
    df.columns = [str(col).strip() for col in df.columns]

    required_columns = ["URL", "Page Type", "Priority", "Notes"]
    missing = [col for col in required_columns if col not in df.columns]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    audit_rows = []
    invalid_rows = []

    print("Auditing URLs...")
    for _, row in df.iterrows():
        url = clean_text(row["URL"])
        page_type = clean_text(row["Page Type"])
        priority = clean_text(row["Priority"])
        notes = clean_text(row["Notes"])

        result = audit_single_url(url, page_type, priority, notes)

        if result["Issues"] in ["Missing URL", "Invalid URL format"] or result["Issues"].startswith("Request failed"):
            invalid_rows.append(result)
        else:
            audit_rows.append(result)

    audit_df = pd.DataFrame(audit_rows)
    invalid_df = pd.DataFrame(invalid_rows)

    summary_df = build_summary(audit_df, invalid_df)

    print("Writing SEO audit workbook...")
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        audit_df.to_excel(writer, sheet_name="SEO_Audit", index=False)
        invalid_df.to_excel(writer, sheet_name="URLs_Invalid", index=False)

    print("Applying Excel formatting...")
    format_workbook(output_file)

    print("\nSEO audit complete.")
    print(f"Created: {output_file.name}")
    print(f"Saved in: {OUTPUTS_DIR}")
    print(f"Audited URLs: {len(audit_df)}")
    print(f"Invalid URL rows: {len(invalid_df)}")


if __name__ == "__main__":
    main()
