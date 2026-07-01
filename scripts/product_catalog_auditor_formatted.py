"""
product_catalog_auditor_formatted.py

Project:
Product Catalog Auditor

Purpose:
Audit a product catalog export, split valid and invalid products, and generate
a formatted Excel report with highlighted issue cells.

Expected folder structure:

Python/
    datasets/
        products_catalog_messy.xlsx
    scripts/
        product_catalog_auditor_formatted.py
    outputs/
    archive/

Run from the Python folder:
    py scripts/product_catalog_auditor_formatted.py

Output format:
    outputs/product_catalog_audit_06.30_1.xlsx
    outputs/product_catalog_audit_06.30_2.xlsx

Output sheets:
    Summary
    Products_Valid
    Products_Invalid

Main features:
    - Cleans product names, categories, prices, stock, status, and dates
    - Validates SKU, product name, category, price, stock, URLs, description, status, and last updated date
    - Moves invalid rows into Products_Invalid
    - Adds final Issues column to Products_Invalid
    - Highlights issue-related cells
    - Sorts invalid rows by SKU so duplicates appear consecutively
    - Formats workbook for portfolio/client presentation
"""

import re
from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent.parent
DATASETS_DIR = BASE_DIR / "datasets"
OUTPUTS_DIR = BASE_DIR / "outputs"
ARCHIVE_DIR = BASE_DIR / "archive"

INPUT_FILE = DATASETS_DIR / "products_catalog_messy.xlsx"

DATASETS_DIR.mkdir(exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)
ARCHIVE_DIR.mkdir(exist_ok=True)


def get_next_output_file(output_dir: Path, base_name: str = "product_catalog_audit", extension: str = ".xlsx") -> Path:
    today = datetime.now().strftime("%m.%d")
    existing_files = list(output_dir.glob(f"{base_name}_{today}_*{extension}"))

    versions = []
    for file in existing_files:
        match = re.search(rf"^{re.escape(base_name)}_{re.escape(today)}_(\d+)$", file.stem)
        if match:
            versions.append(int(match.group(1)))

    next_version = max(versions, default=0) + 1
    return output_dir / f"{base_name}_{today}_{next_version}{extension}"


def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def normalize_category(value):
    return clean_text(value).title()


def normalize_status(value):
    value = clean_text(value).lower()
    mapping = {
        "active": "Active",
        "published": "Active",
        "draft": "Draft",
        "inactive": "Inactive",
    }
    return mapping.get(value, value.title() if value else "")


def clean_price(value):
    if pd.isna(value) or str(value).strip() == "":
        return None

    cleaned = str(value).replace(",", "").strip()

    try:
        return float(cleaned)
    except ValueError:
        return None


def clean_stock(value):
    if pd.isna(value) or str(value).strip() == "":
        return None

    try:
        return int(value)
    except ValueError:
        return None


def parse_date(value):
    if pd.isna(value) or str(value).strip() == "":
        return ""

    parsed = pd.to_datetime(value, errors="coerce", dayfirst=False)

    if pd.isna(parsed):
        parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)

    if pd.isna(parsed):
        return ""

    return parsed.strftime("%Y-%m-%d")


def is_valid_url(value):
    value = clean_text(value)
    return value.startswith("http://") or value.startswith("https://")


def add_issue(issues_by_index, row_index, issue):
    if row_index not in issues_by_index:
        issues_by_index[row_index] = []
    issues_by_index[row_index].append(issue)


def count_total_issues(issues_by_index):
    return sum(len(issues) for issues in issues_by_index.values())


def split_valid_invalid(df, issues_by_index):
    invalid_indexes = sorted(issues_by_index.keys())

    valid_df = df.drop(index=invalid_indexes, errors="ignore").copy()
    invalid_df = df.loc[df.index.intersection(invalid_indexes)].copy()

    if len(invalid_df) > 0:
        invalid_df["Issues"] = invalid_df.index.map(
            lambda idx: "; ".join(sorted(set(issues_by_index.get(idx, []))))
        )

        invalid_df["_original_index"] = invalid_df.index
        invalid_df = invalid_df.sort_values(
            by=["SKU", "_original_index"],
            kind="stable"
        ).drop(columns=["_original_index"])
    else:
        invalid_df["Issues"] = []

    return valid_df.reset_index(drop=True), invalid_df.reset_index(drop=True)


def audit_products(products):
    issues_by_index = {}

    products = products.copy()
    products.columns = [str(col).strip() for col in products.columns]

    required_columns = [
        "SKU", "Product Name", "Category", "Price", "Stock",
        "Image URL", "Product URL", "Description", "Status", "Last Updated"
    ]

    missing_columns = [col for col in required_columns if col not in products.columns]
    if missing_columns:
        raise ValueError(f"Missing required columns: {missing_columns}")

    products["SKU"] = products["SKU"].apply(clean_text)
    products["Product Name"] = products["Product Name"].apply(lambda x: clean_text(x).title())
    products["Category"] = products["Category"].apply(normalize_category)
    products["Price"] = products["Price"].apply(clean_price)
    products["Stock"] = products["Stock"].apply(clean_stock)
    products["Image URL"] = products["Image URL"].apply(clean_text)
    products["Product URL"] = products["Product URL"].apply(clean_text)
    products["Description"] = products["Description"].apply(clean_text)
    products["Status"] = products["Status"].apply(normalize_status)
    products["Last Updated"] = products["Last Updated"].apply(parse_date)

    for index, row in products.iterrows():
        if not row["SKU"]:
            add_issue(issues_by_index, index, "Missing SKU")

        if not row["Product Name"]:
            add_issue(issues_by_index, index, "Missing product name")

        if not row["Category"]:
            add_issue(issues_by_index, index, "Missing category")

        if pd.isna(row["Price"]):
            add_issue(issues_by_index, index, "Missing or invalid price")
        elif row["Price"] <= 0:
            add_issue(issues_by_index, index, "Price is zero or negative")

        if pd.isna(row["Stock"]):
            add_issue(issues_by_index, index, "Missing or invalid stock")
        elif row["Stock"] < 0:
            add_issue(issues_by_index, index, "Stock is negative")
        elif row["Stock"] == 0:
            add_issue(issues_by_index, index, "Out of stock")
        elif row["Stock"] <= 5:
            add_issue(issues_by_index, index, "Low stock")

        if not row["Image URL"]:
            add_issue(issues_by_index, index, "Missing image URL")
        elif not is_valid_url(row["Image URL"]):
            add_issue(issues_by_index, index, "Invalid image URL")

        if not row["Product URL"]:
            add_issue(issues_by_index, index, "Missing product URL")
        elif not is_valid_url(row["Product URL"]):
            add_issue(issues_by_index, index, "Invalid product URL")

        if not row["Description"]:
            add_issue(issues_by_index, index, "Missing description")
        elif len(row["Description"]) < 20:
            add_issue(issues_by_index, index, "Description too short")

        if not row["Status"]:
            add_issue(issues_by_index, index, "Missing status")

        if not row["Last Updated"]:
            add_issue(issues_by_index, index, "Missing or invalid last updated date")

    # Duplicate SKU rule:
    # If a SKU appears more than once, ALL matching rows are invalid.
    duplicate_mask = products.duplicated(subset=["SKU"], keep=False) & products["SKU"].ne("")
    for index in products[duplicate_mask].index:
        add_issue(issues_by_index, index, "Duplicate SKU")

    valid_df, invalid_df = split_valid_invalid(products, issues_by_index)

    return valid_df, invalid_df, issues_by_index


def get_column_index(ws, column_name):
    for cell in ws[1]:
        if cell.value == column_name:
            return cell.column
    return None


def format_workbook(output_file):
    wb = load_workbook(output_file)

    dark_header_fill = PatternFill("solid", fgColor="1F2937")
    dark_header_font = Font(color="FFFFFF", bold=True)
    valid_header_fill = PatternFill("solid", fgColor="166534")
    invalid_header_fill = PatternFill("solid", fgColor="991B1B")
    issue_header_fill = PatternFill("solid", fgColor="D99694")
    issue_fill = PatternFill("solid", fgColor="F4CCCC")
    invalid_cell_fill = PatternFill("solid", fgColor="FCE4E4")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    duplicate_fill = PatternFill("solid", fgColor="FCE5CD")
    summary_fill = PatternFill("solid", fgColor="E0F2FE")

    thin_gray = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    issue_to_columns = {
        "Missing SKU": ["SKU"],
        "Duplicate SKU": ["SKU"],
        "Missing product name": ["Product Name"],
        "Missing category": ["Category"],
        "Missing or invalid price": ["Price"],
        "Price is zero or negative": ["Price"],
        "Missing or invalid stock": ["Stock"],
        "Stock is negative": ["Stock"],
        "Out of stock": ["Stock"],
        "Low stock": ["Stock"],
        "Missing image URL": ["Image URL"],
        "Invalid image URL": ["Image URL"],
        "Missing product URL": ["Product URL"],
        "Invalid product URL": ["Product URL"],
        "Missing description": ["Description"],
        "Description too short": ["Description"],
        "Missing status": ["Status"],
        "Missing or invalid last updated date": ["Last Updated"],
    }

    for ws in wb.worksheets:
        max_row = ws.max_row
        max_col = ws.max_column

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Header formatting
        for cell in ws[1]:
            cell.fill = dark_header_fill
            cell.font = dark_header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        if "Valid" in ws.title:
            for cell in ws[1]:
                cell.fill = valid_header_fill

        if "Invalid" in ws.title:
            for cell in ws[1]:
                cell.fill = invalid_header_fill

        if ws.title == "Summary":
            ws.freeze_panes = None
            for row in ws.iter_rows():
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

            ws["A1"].fill = dark_header_fill
            ws["A1"].font = dark_header_font
            ws["B1"].fill = dark_header_fill
            ws["B1"].font = dark_header_font

            for row in range(2, max_row + 1):
                ws[f"A{row}"].fill = summary_fill

            ws.column_dimensions["A"].width = 34
            ws.column_dimensions["B"].width = 90
            continue

        for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        # Number formats
        price_col = get_column_index(ws, "Price")
        if price_col:
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=price_col).number_format = "#,##0.00"

        stock_col = get_column_index(ws, "Stock")
        if stock_col:
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=stock_col).number_format = "0"

        date_col = get_column_index(ws, "Last Updated")
        if date_col:
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=date_col).number_format = "yyyy-mm-dd"

        # Invalid sheet issue highlighting
        if "Invalid" in ws.title:
            issues_col = get_column_index(ws, "Issues")

            if issues_col:
                ws.cell(row=1, column=issues_col).fill = issue_header_fill
                ws.cell(row=1, column=issues_col).font = Font(color="000000", bold=True)

                for row in range(2, max_row + 1):
                    issues_text = str(ws.cell(row=row, column=issues_col).value or "")
                    ws.cell(row=row, column=issues_col).fill = issue_fill

                    if "Duplicate SKU" in issues_text:
                        for cell in ws[row]:
                            cell.fill = duplicate_fill
                        ws.cell(row=row, column=issues_col).fill = issue_fill

                    if "Low stock" in issues_text:
                        for cell in ws[row]:
                            cell.fill = warning_fill
                        ws.cell(row=row, column=issues_col).fill = issue_fill

                    for issue, columns in issue_to_columns.items():
                        if issue in issues_text:
                            for col_name in columns:
                                col_idx = get_column_index(ws, col_name)
                                if col_idx:
                                    ws.cell(row=row, column=col_idx).fill = invalid_cell_fill

        # Create Excel table
        if max_row >= 2 and max_col >= 1:
            table_name = re.sub(r"[^A-Za-z0-9_]", "", ws.title)[:20] + "Table"
            ref = f"A1:{get_column_letter(max_col)}{max_row}"
            table = Table(displayName=table_name, ref=ref)
            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            try:
                ws.add_table(table)
            except ValueError:
                pass

        # Column widths
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            header = str(ws.cell(row=1, column=col_idx).value or "")

            if header == "Issues":
                ws.column_dimensions[col_letter].width = 58
            elif header in ["Image URL", "Product URL", "Description"]:
                ws.column_dimensions[col_letter].width = 42
            elif header == "Product Name":
                ws.column_dimensions[col_letter].width = 28
            elif header in ["SKU", "Category", "Status", "Last Updated"]:
                ws.column_dimensions[col_letter].width = 18
            else:
                ws.column_dimensions[col_letter].width = 14

        for row in range(1, max_row + 1):
            ws.row_dimensions[row].height = 22

    wb.save(output_file)


def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Input file not found:\n{INPUT_FILE}\n\n"
            "Place products_catalog_messy.xlsx inside the datasets folder."
        )

    output_file = get_next_output_file(OUTPUTS_DIR)

    print("Reading product catalog...")
    print(f"Input: {INPUT_FILE}")

    products = pd.read_excel(INPUT_FILE, sheet_name="Products_Raw")

    print("Auditing products...")
    valid_df, invalid_df, issues = audit_products(products)

    summary = pd.DataFrame([
        ["Input file", str(INPUT_FILE)],
        ["Output file", str(output_file)],
        ["Raw product rows", len(products)],
        ["Products valid rows", len(valid_df)],
        ["Products invalid rows moved", len(invalid_df)],
        ["Product issue count", count_total_issues(issues)],
        ["Duplicate SKU rule", "All rows are invalid when SKU appears more than once"],
        ["Invalid sorting", "Products_Invalid is sorted by SKU so duplicates appear consecutively"],
        ["Formatting", "Issue cells and Issues column are highlighted automatically"],
        ["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ], columns=["Metric", "Value"])

    print("Writing output workbook...")
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        valid_df.to_excel(writer, sheet_name="Products_Valid", index=False)
        invalid_df.to_excel(writer, sheet_name="Products_Invalid", index=False)

    print("Applying Excel formatting...")
    format_workbook(output_file)

    print("\nProduct catalog audit complete.")
    print(f"Created: {output_file.name}")
    print(f"Saved in: {OUTPUTS_DIR}")
    print(f"Products valid rows: {len(valid_df)}")
    print(f"Products invalid rows moved: {len(invalid_df)}")
    print(f"Product issue count: {count_total_issues(issues)}")


if __name__ == "__main__":
    main()
